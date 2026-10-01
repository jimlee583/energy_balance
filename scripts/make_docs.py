"""Generate ``docs/Satellite_Energy_Balance_App.docx``.

Run with::

    uv run --with python-docx python scripts/make_docs.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT_PATH = Path(__file__).resolve().parent.parent / "docs" / "Satellite_Energy_Balance_App.docx"


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)


def _add_para(doc: Document, text: str, style: str | None = None) -> None:
    p = doc.add_paragraph(text, style=style)
    p.paragraph_format.space_after = Pt(6)


def _add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(2)


def _add_code(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    run = p.add_run(text)
    run.font.name = "Menlo"
    r = run._element
    rpr = r.get_or_add_rPr()
    r_fonts = rpr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), "Menlo")
    r_fonts.set(qn("w:hAnsi"), "Menlo")
    run.font.size = Pt(10)


def _add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(header):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row):
            table.rows[r_idx].cells[c_idx].text = val


def build_document() -> Document:
    doc = Document()

    # Narrower margins for a cleaner page.
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Default font tweak.
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # --- Title ----------------------------------------------------------------
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Satellite Energy Balance App")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("How the simulator works, what you can enter, and what it reports")
    run.italic = True
    run.font.size = Pt(12)

    # --- Overview -------------------------------------------------------------
    _add_heading(doc, "1. Overview", 1)
    _add_para(
        doc,
        "This application answers a basic but critical question for every satellite mission: "
        "does the solar array produce enough power, over the course of each orbit and across the "
        "seasons, to keep the loads running and the battery charged? The user describes the "
        "orbit, the attitude, the solar panels, the loads, and the battery; the app propagates "
        "the orbit over time, tracks sunlight and eclipse, computes panel power, integrates the "
        "battery state of charge, and reports whether the energy budget closes.",
    )
    _add_para(
        doc,
        "The app has two layers. A plain Python physics package (energy_balance/) performs all "
        "calculations and has no dependency on the user interface, so it can be reused from "
        "notebooks, scripts, or tests. A thin Streamlit layer (app.py) wraps the package with "
        "input forms, interactive Plotly charts, and CSV exports.",
    )

    _add_heading(doc, "1.1 How to install and run", 2)
    _add_code(
        doc,
        "uv sync --all-extras\n"
        "uv run streamlit run app.py   # start the UI\n"
        "uv run pytest                 # run validation tests",
    )

    # --- Architecture ---------------------------------------------------------
    _add_heading(doc, "2. Architecture at a glance", 1)
    _add_para(
        doc,
        "Each simulation step runs through the same pipeline. The orbit propagator advances the "
        "satellite's position; the solar ephemeris places the Sun; the eclipse model combines the "
        "two to produce an illumination fraction between 0 and 1; the attitude model builds a "
        "body-to-ECI rotation matrix; the solar array uses that matrix to compute per-panel "
        "power; the loads are evaluated against the illumination fraction; and the battery "
        "integrator closes the loop.",
    )
    _add_table(
        doc,
        header=["Stage", "Module", "What it produces"],
        rows=[
            ["Inputs", "config.py", "Validated SimulationConfig (pydantic)"],
            ["Orbit", "orbit.py", "ECI position / velocity at every time step"],
            ["Sun", "sun.py", "ECI direction and distance of the Sun"],
            ["Eclipse", "eclipse.py", "Illumination fraction 0..1"],
            ["Attitude", "attitude.py", "Rotation matrix body -> ECI"],
            ["Panels", "solar_array.py", "Per-panel generated power (W)"],
            ["Loads", "loads.py", "Per-load demand (W)"],
            ["Battery", "battery.py", "State of charge, shunt, unmet load"],
            ["Pipeline", "simulation.py", "Tidy pandas DataFrame of everything above"],
            ["Reporting", "metrics.py", "Verdict, per-orbit table, sizing hints"],
            ["UI", "app.py", "Sidebar, tabs, plots, downloads"],
        ],
    )

    # --- User inputs ----------------------------------------------------------
    _add_heading(doc, "3. What the user enters", 1)
    _add_para(
        doc,
        "All inputs are collected in the Streamlit sidebar. There are three built-in presets "
        "(an ISS-like LEO orbit, a 550 km sun-synchronous orbit at 10:30 LTAN, and a GEO orbit); "
        "the user can start from any of them, tweak values, and save or load the whole config as "
        "a JSON file.",
    )

    _add_heading(doc, "3.1 Mission", 2)
    _add_bullets(
        doc,
        [
            "Epoch (UTC date and time) - fixes the Sun's position and the season.",
            "Simulation duration in days.",
            "Time step in seconds (default 30 s; 120 s for GEO).",
        ],
    )

    _add_heading(doc, "3.2 Orbit", 2)
    _add_bullets(
        doc,
        [
            "Specify the orbit by perigee / apogee altitudes OR by semi-major axis + eccentricity.",
            "Inclination (deg), RAAN (deg), argument of perigee (deg), true anomaly (deg).",
            "Optional 'sun-synchronous' flag: the app computes the inclination automatically from the altitude, and places RAAN according to the chosen local time of the ascending node (LTAN).",
        ],
    )

    _add_heading(doc, "3.3 Attitude", 2)
    _add_bullets(
        doc,
        [
            "Nadir-pointing (+Z body axis to Earth, chosen body axis along velocity).",
            "Sun-pointing (chosen body axis to the Sun).",
            "Inertially fixed (body axes aligned with ECI).",
        ],
    )

    _add_heading(doc, "3.4 Solar arrays", 2)
    _add_para(
        doc,
        "The user adds one or more panels. Each panel has:",
    )
    _add_bullets(
        doc,
        [
            "Area (m^2) and a body-frame normal vector (named face like +X / -Y, or a custom direction).",
            "Mounting: fixed, 1-axis tracking (about a body-frame axis), or 2-axis tracking (always faces the Sun).",
            "Cell efficiency, packing factor, inherent degradation, degradation per year, mission year (which combine to give an end-of-life factor), and a temperature loss factor.",
            "A global solar constant (default 1361 W/m^2); the app scales this by the Earth-Sun distance.",
        ],
    )

    _add_heading(doc, "3.5 Loads", 2)
    _add_para(
        doc,
        "Loads are entered in an editable table. Each row has a name, a power in watts, and a mode "
        "that controls when it runs:",
    )
    _add_bullets(
        doc,
        [
            "always - constant power throughout the orbit.",
            "sunlit_only - active only when the satellite is in sunlight (e.g. an imager).",
            "eclipse_only - active only in shadow (e.g. survival heaters).",
            "duty_cycle - a constant fraction of the orbit (duty_percent).",
        ],
    )

    _add_heading(doc, "3.6 Battery and power system", 2)
    _add_bullets(
        doc,
        [
            "Battery: capacity (Wh), initial state of charge (%), maximum depth of discharge (%), charge and discharge efficiencies, and maximum charge rate (W).",
            "Power system: regulation type (MPPT or Direct Energy Transfer) and the end-to-end path efficiency.",
        ],
    )

    # --- Physics models -------------------------------------------------------
    _add_heading(doc, "4. How the physics is computed", 1)

    _add_heading(doc, "4.1 Orbit propagation (orbit.py)", 2)
    _add_para(
        doc,
        "The orbit is held as classical Keplerian elements. At every time step the mean anomaly is "
        "advanced, Kepler's equation is solved for the eccentric anomaly, and the position and "
        "velocity are rotated out of the perifocal frame into the Earth-centered inertial (ECI) "
        "frame. To capture the dominant long-term perturbation, the right ascension of the "
        "ascending node (RAAN) and the argument of perigee are drifted linearly with the secular "
        "J2 rates, and the mean motion gets the J2 correction:",
    )
    _add_code(
        doc,
        "RAAN_dot  = -3/2 * n * J2 * (R_e / p)^2 * cos(i)\n"
        "argp_dot  =  3/4 * n * J2 * (R_e / p)^2 * (5 cos^2(i) - 1)\n"
        "M_dot_corr = n + 3/4 * n * J2 * (R_e/p)^2 * sqrt(1-e^2) * (3 cos^2(i) - 1)",
    )
    _add_para(
        doc,
        "The sun-synchronous helper inverts the RAAN equation to pick the inclination whose RAAN "
        "drift exactly matches Earth's mean orbital rate around the Sun (about 0.9856 deg/day). "
        "The LTAN-to-RAAN helper uses the mean Sun's right ascension at the epoch and offsets it "
        "by (LTAN - 12) * 15 deg.",
    )

    _add_heading(doc, "4.2 Solar ephemeris (sun.py)", 2)
    _add_para(
        doc,
        "The Sun's position is computed with the low-precision formula from the Astronomical "
        "Almanac / Vallado, which is accurate to roughly 0.01 degree for dates within a few "
        "decades of J2000 - far more than enough for an eclipse analysis. The output is a unit "
        "vector in ECI and the Earth-Sun distance in AU (which scales the solar flux).",
    )

    _add_heading(doc, "4.3 Eclipse model (eclipse.py)", 2)
    _add_para(
        doc,
        "The conical shadow model computes the apparent angular radii of the Sun and Earth as "
        "seen from the satellite and their angular separation; the illumination fraction is 1 "
        "minus the fraction of the Sun's disk hidden by the Earth. Full umbra gives 0, full "
        "sunlight gives 1, and the penumbra crossing is smooth. A simpler cylindrical shadow "
        "model is also available and is used for cross-checking in tests.",
    )

    _add_heading(doc, "4.4 Attitude (attitude.py)", 2)
    _add_para(
        doc,
        "For each time step the app builds a 3x3 body-to-ECI rotation matrix:",
    )
    _add_bullets(
        doc,
        [
            "Nadir pointing: +Z body axis is placed along the nadir direction, and the chosen yaw axis is placed along the velocity direction.",
            "Sun pointing: the chosen body axis is placed along the Sun direction; the remaining axes complete the right-handed frame.",
            "Inertial: the rotation is the identity (body == ECI).",
        ],
    )

    _add_heading(doc, "4.5 Solar array power (solar_array.py)", 2)
    _add_para(
        doc,
        "For each panel, the instantaneous power is:",
    )
    _add_code(
        doc,
        "P = flux * area * efficiency * packing * I_d * (1 - deg)^year *\n"
        "    (1 - temp_loss) * max(cos(incidence_angle), 0) * illumination",
    )
    _add_para(
        doc,
        "The incidence angle is the angle between the panel's normal and the Sun direction. For "
        "fixed panels the normal is rotated into ECI with the attitude matrix. For 1-axis tracking "
        "panels the app projects the Sun direction onto the plane perpendicular to the rotation "
        "axis and takes the magnitude - this is the best achievable cosine. For 2-axis tracking "
        "the cosine is always 1 while the satellite is in sunlight.",
    )

    _add_heading(doc, "4.6 Loads (loads.py)", 2)
    _add_para(
        doc,
        "Each load is turned into a per-step power demand based on its mode. 'Sunlit' loads are "
        "multiplied by the illumination fraction so they ramp down smoothly through the penumbra; "
        "'eclipse' loads are multiplied by (1 - illumination).",
    )

    _add_heading(doc, "4.7 Battery (battery.py)", 2)
    _add_para(
        doc,
        "The battery is a leaky integrator. The net power going into the battery is the generated "
        "power times the path efficiency, minus the load. When positive, the charge rate is "
        "limited to the configured maximum and to whatever room is left in the battery; the "
        "charge efficiency is applied; excess is reported as 'shunt' (wasted) power. When "
        "negative, the deficit is pulled from the battery at the discharge efficiency; if the "
        "battery is empty, the shortfall is reported as unmet load.",
    )

    # --- Outputs --------------------------------------------------------------
    _add_heading(doc, "5. What the app shows", 1)
    _add_para(
        doc,
        "The UI has five tabs. The sidebar also has a 'Download results (CSV)' button that "
        "exports the entire time-series.",
    )

    _add_heading(doc, "5.1 Summary", 2)
    _add_bullets(
        doc,
        [
            "Orbit period, inclination, mean eclipse duration, and RAAN drift per day.",
            "Totals for energy generated, consumed, shunted, and unmet.",
            "Minimum state of charge, maximum observed depth of discharge, and beta-angle range.",
            "A green PASS or red FAIL verdict, with four individual checks: power budget positive, DoD within limit, no unmet load, and battery recovers each orbit.",
            "A per-orbit energy table and sizing hints (minimum solar array area and minimum battery capacity to close the budget at the worst-case orbit).",
        ],
    )

    _add_heading(doc, "5.2 Power", 2)
    _add_bullets(
        doc,
        [
            "A stacked time-series plot of generated, load, and net power, with an illumination strip below.",
            "An expandable per-panel breakdown and an expandable per-load breakdown.",
        ],
    )

    _add_heading(doc, "5.3 Battery", 2)
    _add_bullets(
        doc,
        [
            "State-of-charge time series with the depth-of-discharge limit drawn as a red dashed line.",
            "Shunted and unmet power stacked below.",
            "Beta angle time series for context.",
        ],
    )

    _add_heading(doc, "5.4 Orbit", 2)
    _add_bullets(
        doc,
        [
            "A 3D view of one revolution around the Earth with the Sun direction marked.",
            "A ground track for the first few orbits (computed with Greenwich mean sidereal time to approximate ECI -> ECEF).",
        ],
    )

    _add_heading(doc, "5.5 Seasonal", 2)
    _add_para(
        doc,
        "A one-year sweep of the beta angle and the analytic eclipse fraction, computed from the "
        "orbit's angular momentum direction and the Sun's position every five days. This is the "
        "fastest way to find the worst-case eclipse season for a long-term mission.",
    )

    # --- Validation -----------------------------------------------------------
    _add_heading(doc, "6. How we know it works", 1)
    _add_para(
        doc,
        "The test suite (tests/) contains 26 validation checks against textbook values and "
        "analytic formulas. Every one of them must pass before the app is considered correct.",
    )
    _add_table(
        doc,
        header=["Expectation", "Result in the test suite"],
        rows=[
            ["ISS-like orbit period at 420 km", "~92.96 min (textbook ~92.7 min)"],
            ["Sun-synchronous inclination at 550 km", "97.6 deg"],
            ["Sun-synchronous inclination at 800 km", "98.6 deg"],
            ["Sun-synchronous RAAN drift", "0.9856 deg/day (matches Earth's orbit rate)"],
            ["Analytic eclipse at 400 km, beta = 0", "~36 min per orbit"],
            ["Perigee/apogee radii of an eccentric orbit", "Match input altitudes to within 1 km"],
            ["Circular-orbit speed", "Matches vis-viva"],
            ["Battery state of charge", "Always in [0, 100] %"],
            ["Energy conservation in a closed loop", "Delta SOC = net energy in"],
            ["Config save/load round trip", "Byte-for-byte identical"],
            ["All three presets", "Run end-to-end without error"],
        ],
    )

    # --- Extensibility --------------------------------------------------------
    _add_heading(doc, "7. Extending the app", 1)
    _add_para(
        doc,
        "Because the physics modules are decoupled from the UI and from each other, each one "
        "can be swapped independently. The most likely future upgrades are:",
    )
    _add_bullets(
        doc,
        [
            "Importing orbits from two-line element sets (TLEs) using the SGP4 propagator.",
            "Adding Earth albedo and infrared re-radiation as additional sunlight sources.",
            "Modeling panel self-shadowing by the spacecraft bus or by other panels.",
            "Simulating cell temperature as a function of orbit position and dissipating power.",
            "Running Monte Carlo cases over attitude errors or component tolerances.",
        ],
    )

    # --- Appendix: input reference -------------------------------------------
    _add_heading(doc, "Appendix A. Input reference", 1)
    _add_para(
        doc,
        "The pydantic models in energy_balance/config.py are the canonical source of input "
        "documentation; the table below is a quick reference.",
    )
    _add_table(
        doc,
        header=["Config class", "Key fields"],
        rows=[
            [
                "MissionConfig",
                "epoch, duration_days, time_step_s",
            ],
            [
                "OrbitConfig",
                "spec, perigee/apogee altitudes OR semi-major axis + eccentricity, inclination, RAAN, arg_perigee, true_anomaly, sun_synchronous, ltan_hours",
            ],
            [
                "AttitudeConfig",
                "mode (nadir / sun / inertial), nadir_yaw_axis, sun_pointing_axis",
            ],
            [
                "PanelConfig",
                "name, area_m2, normal_face / normal_body, mounting (fixed / 1-axis / 2-axis), rotation_axis_body, cell_efficiency, packing_factor, inherent_degradation, degradation_per_year, mission_year, temperature_loss",
            ],
            [
                "SolarArrayConfig",
                "panels (list of PanelConfig), solar_constant_w_m2",
            ],
            [
                "LoadConfig",
                "name, power_w, mode (always / sunlit / eclipse / duty_cycle), duty_percent",
            ],
            [
                "BatteryConfig",
                "capacity_wh, initial_soc_percent, max_dod_percent, charge_efficiency, discharge_efficiency, max_charge_rate_w",
            ],
            [
                "PowerSystemConfig",
                "regulation (MPPT / DET), path_efficiency",
            ],
            [
                "SimulationConfig",
                "name, mission, orbit, attitude, solar_array, loads, battery, power_system",
            ],
        ],
    )

    return doc


def main() -> None:
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
